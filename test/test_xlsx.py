# -*- coding: utf-8 -*-
# SPDX-FileCopyrightText: 2026 RYXLAB SOFT TECH SRL (RyxLabs)
# SPDX-License-Identifier: GPL-2.0-or-later
"""The Excel writer and reader. Files are read back by CADex's reader and, when installed,
by openpyxl. No QgsApplication."""

import importlib.util
import os
import shutil
import tempfile
import unittest
import zipfile

import helpers  # noqa: F401  sys.path bootstrap

from cadex.services import diminuare_service, xlsx, xlsx_report

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None


def load_with_openpyxl(path):
    # openpyxl 3.1.2, the version in QGIS 3.40, leaves a workbook it was given by name open
    import openpyxl
    with open(path, "rb") as fp:
        return openpyxl.load_workbook(fp)


class WriterTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_xlsx_")
        self.path = os.path.join(self.tmp, "raport.xlsx")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self):
        wb = xlsx_report.new_workbook()
        rows = [["10188", 2628.0, "POPESCU ION", None],
                ["0123", 1.5, "text cu <tag> & \"ghilimele\"", "rând 1\nrând 2"],
                ["x", -3, "  spații  ", "ș ț ă î â Ş Ţ"]]
        xlsx_report.write_table(wb, "Imobile / sector 1", ["id", "sup", "propr", "note"],
                                rows, row_colors=[xlsx_report.GOOD, None, None],
                                text_columns=["id"])
        summary = xlsx_report.Summary(wb, "Titlu")
        summary.section("Informații")
        summary.facts([("Data", "24.09.2026"), ("Fișiere", 3)])
        summary.table(["Categorie", "Nr."], [["A", 1200], ["B", 3.5]])
        summary.note("o notă")
        xlsx_report.save(wb, self.path)

    def test_round_trip_through_our_reader(self):
        self.build()
        self.assertEqual(xlsx.sheet_names(self.path), ["Sumar", "Imobile - sector 1"])
        _names, rows = xlsx.read_rows(self.path, "Imobile - sector 1")
        self.assertEqual(rows[0], ["id", "sup", "propr", "note"])
        self.assertEqual(rows[1], ["10188", 2628.0, "POPESCU ION", None])
        self.assertEqual(rows[2][0], "0123")  # stored as text, the leading zero stays
        self.assertEqual(rows[2][2], 'text cu <tag> & "ghilimele"')
        self.assertEqual(rows[2][3], "rând 1\nrând 2")
        self.assertEqual(rows[3][2], "  spații  ")
        self.assertEqual(rows[3][3], "ș ț ă î â Ş Ţ")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl missing")
    def test_openpyxl_reads_structure_and_styles(self):
        self.build()
        wb = load_with_openpyxl(self.path)
        self.assertEqual(wb.sheetnames, ["Sumar", "Imobile - sector 1"])
        ws = wb["Imobile - sector 1"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:D4")
        header = ws["A1"]
        self.assertTrue(header.font.b)
        self.assertEqual(header.fill.fgColor.rgb, "FF" + xlsx_report.ACCENT)
        self.assertEqual(ws["A2"].fill.fgColor.rgb, "FF" + xlsx_report.GOOD)
        self.assertEqual(ws["A3"].number_format, "@")
        self.assertEqual(ws["B2"].value, 2628)
        summary = wb["Sumar"]
        self.assertIn("A1:D1", [str(r) for r in summary.merged_cells.ranges])
        self.assertFalse(summary.sheet_view.showGridLines)
        self.assertEqual(summary["B8"].value, 1200)
        self.assertEqual(summary["B8"].number_format, "#,##0")

    def test_percentages_are_numbers_and_text_numbers_read_romanian(self):
        wb = xlsx_report.new_workbook()
        summary = xlsx_report.Summary(wb, "Titlu")
        summary.table(["Situație", "Procent"], [["Intravilan", xlsx_report.percent(55.555, 1)]])
        # old and new values of one column can be text in one row and numbers in the next
        xlsx_report.write_table(wb, "Modificate", ["Ce", "Valoare"],
                                [["sup", 1155.0], ["sup", 1186.5], ["nr", "0123"]],
                                text_columns=["Valoare"])
        xlsx_report.save(wb, self.path)
        _names, rows = xlsx.read_rows(self.path, "Modificate")
        self.assertEqual([r[1] for r in rows[1:]], ["1155", "1186,5", "0123"])
        _names, rows = xlsx.read_rows(self.path, "Sumar")
        # a number Excel shows as 55,6% or 55.6% after the computer's settings
        self.assertAlmostEqual(rows[3][1], 0.55555)
        if HAS_OPENPYXL:
            self.assertEqual(load_with_openpyxl(self.path)["Sumar"]["B4"].number_format,
                             "0.0%")

    def test_package_parts(self):
        self.build()
        with zipfile.ZipFile(self.path) as zf:
            names = set(zf.namelist())
            workbook = zf.read("xl/workbook.xml").decode("utf-8")
        self.assertTrue({"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                         "xl/_rels/workbook.xml.rels", "xl/styles.xml",
                         "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"} <= names)
        # a filter without its hidden defined name makes Excel "repair" the file
        self.assertIn("_xlnm._FilterDatabase", workbook)
        self.assertIn("'Imobile - sector 1'!$A$1:$D$4", workbook)

    def test_sheet_names_are_made_legal_and_unique(self):
        wb = xlsx.Workbook()
        self.assertEqual(wb.add_sheet("a/b:c?").name, "a-b-c-")
        self.assertEqual(wb.add_sheet("x" * 40).name, "x" * 31)
        self.assertEqual(wb.add_sheet("A-B-C-").name, "A-B-C- (2)")

    def test_empty_workbook_still_opens(self):
        xlsx.Workbook().save(self.path)
        self.assertEqual(xlsx.sheet_names(self.path), ["Foaie"])

    def test_column_letters(self):
        for index, letters in ((1, "A"), (26, "Z"), (27, "AA"), (702, "ZZ"), (703, "AAA")):
            self.assertEqual(xlsx.column_letter(index), letters)
            self.assertEqual(xlsx.column_index(letters), index)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl missing to write the fixtures")
class ReaderTests(unittest.TestCase):
    """Files written by Excel-like producers: shared strings, formulas with cached values, gaps."""

    def setUp(self):
        import openpyxl
        self.tmp = tempfile.mkdtemp(prefix="cadex_read_")
        self.path = os.path.join(self.tmp, "parcele.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Parcele"
        ws.append(["Nr", "Suprafata act", "IE"])
        ws.append([1, 1500, None])
        ws.append([])
        ws.append([2, "1.234,5", "50124"])
        ws["D6"] = "departe"
        wb.create_sheet("Altceva").append(["x"])
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_gaps_and_row_numbers(self):
        names, rows = xlsx.read_rows(self.path)
        self.assertEqual(names, ["Parcele", "Altceva"])
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[2], [None, None, None, None])
        self.assertEqual(rows[5][3], "departe")
        _names, header, data, numbers, header_row, _hidden = diminuare_service.read_sheet(
            self.path)
        self.assertEqual(header_row, 1)
        # the single value in column D gets its own column, with the generated header "Coloana 4"
        self.assertEqual(header, ["Nr", "Suprafata act", "IE", "Coloana 4"])
        # the blank row 3 is skipped but the numbers still match what Excel shows
        self.assertEqual(numbers, [2, 4, 6])
        self.assertEqual(data[1], [2.0, "1.234,5", "50124", None])

    def test_other_sheet(self):
        _names, rows = xlsx.read_rows(self.path, "Altceva")
        self.assertEqual(rows, [["x"]])

    def test_old_xls_is_refused_clearly(self):
        fake = os.path.join(self.tmp, "vechi.xlsx")
        with open(fake, "wb") as fp:
            fp.write(b"\xd0\xcf\x11\xe0 not a zip")
        with self.assertRaises(zipfile.BadZipFile):
            xlsx.read_rows(fake)


class CsvTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_csv_")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name, data: bytes):
        path = os.path.join(self.tmp, name)
        with open(path, "wb") as fp:
            fp.write(data)
        return path

    def test_romanian_excel_csv(self):
        # Excel with Romanian regional settings: ";" between columns, "," for decimals, code
        # page 1250, which only has the cedilla ţ
        path = self.write("ro.csv",
                          "Nr;Suprafaţă act;Proprietar\n1;1500,5;ŞTEFAN\n".encode("cp1250"))
        _names, rows = xlsx.read_rows(path)
        self.assertEqual(rows[0], ["Nr", "Suprafaţă act", "Proprietar"])
        self.assertEqual(rows[1], ["1", "1500,5", "ŞTEFAN"])
        self.assertEqual(diminuare_service.parse_number(rows[1][1]), 1500.5)

    def test_utf8_bom_comma_csv(self):
        path = self.write("en.csv", "﻿Nr,Area\n1,1500.5\n".encode("utf-8"))
        _names, rows = xlsx.read_rows(path)
        self.assertEqual(rows, [["Nr", "Area"], ["1", "1500.5"]])

    def test_empty_cells_are_none(self):
        path = self.write("gol.csv", b"a;b;c\n1;;3\n")
        self.assertEqual(xlsx.read_rows(path)[1][1], ["1", None, "3"])


def sheet_xml(path, number=1):
    with zipfile.ZipFile(path) as zf:
        return zf.read(f"xl/worksheets/sheet{number}.xml").decode("utf-8")


def cell_format(path, ref, number=1):
    """(number format code, wraps) of a cell, read from the package parts."""
    import re
    # parses only a file this test wrote
    import xml.etree.ElementTree as ET  # nosec B405
    ns = {"m": xlsx.MAIN_NS}
    match = re.search(rf'<c r="{ref}"(?: s="(\d+)")?', sheet_xml(path, number))
    with zipfile.ZipFile(path) as zf:
        styles = ET.fromstring(zf.read("xl/styles.xml"))  # nosec B314
    xf = styles.find("m:cellXfs", ns)[int(match.group(1) or 0)]
    codes = {int(f.get("numFmtId")): f.get("formatCode")
             for f in styles.iter(f"{{{xlsx.MAIN_NS}}}numFmt")}
    codes.update({number: code for code, number in xlsx.BUILTIN_FORMATS.items()})
    alignment = xf.find("m:alignment", ns)
    wraps = alignment is not None and alignment.get("wrapText") == "1"
    return codes.get(int(xf.get("numFmtId"))), wraps


def row_height(path, row, number=1):
    import re
    match = re.search(rf'<row r="{row}"(?: ht="([\d.]+)")?', sheet_xml(path, number))
    return float(match.group(1)) if match and match.group(1) else None


class ServicesWorkbookSaveTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_save_")
        self.path = os.path.join(self.tmp, "raport.xlsx")
        wb = xlsx.Workbook()
        wb.add_sheet("Raportul de ieri")
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_a_save_that_fails_midway_keeps_the_previous_workbook(self):
        wb = xlsx.Workbook()
        sheet = wb.add_sheet("Raportul de azi")
        sheet.set(1, 1, "x")

        def disk_full():
            raise OSError(28, "No space left on device")
        sheet.xml = disk_full
        with self.assertRaises(OSError) as raised:
            wb.save(self.path)
        # the message names the workbook the user chose, not a temporary file
        self.assertEqual(raised.exception.filename, self.path)
        self.assertEqual(xlsx.sheet_names(self.path), ["Raportul de ieri"])
        self.assertEqual(os.listdir(self.tmp), ["raport.xlsx"])

    def test_a_save_replaces_the_previous_workbook_and_leaves_no_other_file(self):
        wb = xlsx.Workbook()
        wb.add_sheet("Raportul de azi")
        wb.save(self.path)
        self.assertEqual(xlsx.sheet_names(self.path), ["Raportul de azi"])
        self.assertEqual(os.listdir(self.tmp), ["raport.xlsx"])

    @unittest.skipUnless(os.name == "nt", "Excel's way of holding a file is a Windows lock")
    def test_a_workbook_open_in_excel_is_refused_by_name_and_kept(self):
        from test_pure_services import hold_like_excel
        release = hold_like_excel(self.path)
        try:
            wb = xlsx.Workbook()
            wb.add_sheet("Raportul de azi")
            with self.assertRaises(PermissionError) as raised:
                wb.save(self.path)
        finally:
            release()
        self.assertEqual(raised.exception.filename, self.path)
        self.assertEqual(xlsx.sheet_names(self.path), ["Raportul de ieri"])
        self.assertEqual(os.listdir(self.tmp), ["raport.xlsx"])

    def test_a_missing_folder_is_reported_for_the_workbook(self):
        path = os.path.join(self.tmp, "lipsa", "raport.xlsx")
        with self.assertRaises(FileNotFoundError) as raised:
            xlsx.Workbook().save(path)
        self.assertEqual(raised.exception.filename, path)


class ServicesReaderTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_hidden_")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_rows_hidden_in_excel_are_reported_by_number(self):
        path = os.path.join(self.tmp, "parcele.xlsx")
        wb = xlsx.Workbook()
        sheet = wb.add_sheet("Parcele")
        for row, values in enumerate([["Nr", "Tarla"], [1, 12], [2, 15], [3, 12]], start=1):
            for col, value in enumerate(values, start=1):
                sheet.set(row, col, value)
        wb.save(path)
        # a filter on tarla 12 hides row 3, the way Excel saves it
        with zipfile.ZipFile(path) as zf:
            parts = {name: zf.read(name) for name in zf.namelist()}
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
            b'<row r="3"', b'<row r="3" hidden="1"')
        with zipfile.ZipFile(path, "w") as zf:
            for name, data in parts.items():
                zf.writestr(name, data)
        names, rows, hidden = xlsx.read_table(path)
        self.assertEqual(names, ["Parcele"])
        # every row is still read, so row numbers match Excel's
        self.assertEqual([r[0] for r in rows], ["Nr", 1.0, 2.0, 3.0])
        self.assertEqual(hidden, [3])
        self.assertEqual(xlsx.read_rows(path), (names, rows))

    def test_a_csv_has_no_hidden_rows(self):
        path = os.path.join(self.tmp, "parcele.csv")
        with open(path, "wb") as fp:
            fp.write(b"Nr;Tarla\n1;12\n")
        self.assertEqual(xlsx.read_table(path), (["CSV"], [["Nr", "Tarla"], ["1", "12"]], []))

    def test_a_romanian_csv_with_a_comma_in_its_header_keeps_its_columns(self):
        path = os.path.join(self.tmp, "parcele.csv")
        with open(path, "wb") as fp:
            fp.write("Nr;Suprafata, mp;IE\n1;1001,5;\n2;1002,5;50124\n".encode("cp1250"))
        _names, rows = xlsx.read_rows(path)
        self.assertEqual(rows, [["Nr", "Suprafata, mp", "IE"], ["1", "1001,5", None],
                                ["2", "1002,5", "50124"]])


class ServicesSummaryTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_sumar_")
        self.path = os.path.join(self.tmp, "raport.xlsx")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_long_facts_and_notes_are_shown_whole(self):
        columns = ", ".join(f"Coloana numărul {i}" for i in range(1, 14))
        note = " ".join(["Imobilele fără coordonate în exportul eTerra nu se compară."] * 9)
        wb = xlsx_report.new_workbook()
        summary = xlsx_report.Summary(wb, "Comparator")
        summary.facts([("Data", "24.09.2026"), ("Coloane comparate", columns)])
        summary.note(note)
        xlsx_report.save(wb, self.path)
        self.assertGreater(len(columns), 240)
        self.assertGreater(len(note), 500)
        # facts at rows 3 and 4, the note at row 6
        self.assertTrue(cell_format(self.path, "B4")[1])
        self.assertIsNone(row_height(self.path, 3))
        self.assertGreaterEqual(row_height(self.path, 4), 3 * 15)
        self.assertGreaterEqual(row_height(self.path, 6), 3 * 12)

    def test_the_summary_prints_one_page_wide(self):
        wb = xlsx_report.new_workbook()
        xlsx_report.Summary(wb, "Titlu").note("o notă")
        xlsx_report.write_table(wb, "Date", ["a"], [[1]])
        xlsx_report.save(wb, self.path)
        summary = sheet_xml(self.path, 1)
        # the schema orders sheetPr first and pageSetup right after pageMargins
        self.assertRegex(summary, r'<worksheet [^>]*><sheetPr><pageSetUpPr fitToPage="1"/>')
        self.assertRegex(summary, r'<pageMargins [^>]*/><pageSetup paperSize="9" '
                                  r'orientation="portrait" fitToWidth="1" fitToHeight="0"/>')
        self.assertNotIn("pageSetup", sheet_xml(self.path, 2))
        if HAS_OPENPYXL:
            ws = load_with_openpyxl(self.path)["Sumar"]
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 0))


class ServicesCellValueTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cadex_dates_")
        self.path = os.path.join(self.tmp, "raport.xlsx")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_dates_are_excel_dates_that_sort_by_time(self):
        from datetime import date, datetime
        wb = xlsx_report.new_workbook()
        xlsx_report.write_table(wb, "Date", ["Data", "Măsurat"], [
            [date(2025, 3, 1), datetime(2024, 1, 2, 14, 30)],
            [date(2024, 1, 2), None]])
        xlsx_report.save(wb, self.path)
        _names, rows = xlsx.read_rows(self.path)
        # day numbers: 2 January 2024 is 45293 and comes before 1 March 2025
        self.assertEqual([r[0] for r in rows[1:]], [45717.0, 45293.0])
        self.assertAlmostEqual(rows[1][1], 45293 + 14.5 / 24)
        self.assertEqual(cell_format(self.path, "A2")[0], "dd.mm.yyyy")
        self.assertEqual(cell_format(self.path, "B2")[0], "dd.mm.yyyy hh:mm")
        if HAS_OPENPYXL:
            ws = load_with_openpyxl(self.path)["Date"]
            self.assertEqual(ws["A3"].value, datetime(2024, 1, 2))

    def test_times_are_written_as_times_and_qt_dates_as_dates(self):
        from datetime import time
        self.assertEqual(xlsx_report.cell_value(time(14, 30, 5)), "14:30:05")
        try:
            from qgis.PyQt.QtCore import QDate, QTime  # type: ignore
        except ImportError:
            self.skipTest("Qt types need QGIS's Python")
        self.assertEqual(xlsx_report.cell_value(QTime(14, 30, 5)), "14:30:05")
        self.assertEqual(xlsx_report.cell_value(QDate(2024, 1, 2)).value, 45293)


if __name__ == "__main__":
    unittest.main()
